"""Deliverable workbook: editable model tabs (live formulas), pricing tab (DK price inputs -> EV/Kelly), raw tables."""
import os, sys, json, numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
OUT = os.path.join(os.path.dirname(__file__), '..', 'output')
res = json.load(open(f'{OUT}/qf_results.json')); mk = json.load(open(f'{OUT}/markets.json')); card = json.load(open(f'{OUT}/card.json'))
REC = json.load(open(f'{OUT}/reconcile.json')); XC = json.load(open(f'{OUT}/crosscheck.json'))
BLUE = Font(color='0000CC', bold=True); HDR = Font(bold=True); TITLE = Font(bold=True, size=13)
FILL = PatternFill('solid', fgColor='EEF3FB')

def hold_f(p):
    q = f'(1-{p})'
    return f'({p}^4*(1+4*{q}+10*{q}^2)+20*{p}^3*{q}^3*{p}^2/(1-2*{p}*{q}))'

def tb_block(ws, r0, c0, pa, pb):
    """7x7 DP for a TB with A serving first; returns cell with P(A wins TB)."""
    def cell(i, j): return ws.cell(r0 + i, c0 + j).coordinate
    def a_pt(k):  # k = 0-based point index
        first = (k == 0) or (((k - 1) // 2) % 2 == 1)
        return pa if first else f'(1-{pb})'
    for i in range(7):
        for j in range(7):
            if i == 0 and j == 0: ws.cell(r0, c0, 1); continue
            t = []
            if i > 0: t.append(f'{cell(i-1, j)}*{a_pt(i - 1 + j)}')
            if j > 0: t.append(f'{cell(i, j-1)}*(1-{a_pt(i + j - 1)})')
            ws.cell(r0 + i, c0 + j, '=' + '+'.join(t))
    win = [f'{cell(6, j)}*{a_pt(6 + j)}' for j in range(6)]
    w = f'({pa}*(1-{pb}))'; l = f'((1-{pa})*{pb})'
    ws.cell(r0 + 7, c0 - 1, 'P(A wins TB)')
    ws.cell(r0 + 7, c0, '=' + '+'.join(win) + f'+{cell(6, 6)}*{w}/({w}+{l})')
    return ws.cell(r0 + 7, c0).coordinate

def set_block(ws, r0, c0, hA, hB, tbc):
    def cell(i, j): return ws.cell(r0 + i, c0 + j).coordinate
    def a_game(g):  # g = 1-based game number; A serves odd games
        return hA if g % 2 == 1 else f'(1-{hB})'
    for i in range(7):
        for j in range(7):
            if i == 0 and j == 0: ws.cell(r0, c0, 1); continue
            t = []
            if i > 0 and not ((i - 1) >= 6 and (i - 1) - j >= 2) and not (j >= 6 and j - (i - 1) >= 2) and not (i - 1 == 6 and j < 5):
                t.append(f'{cell(i-1, j)}*{a_game(i - 1 + j + 1)}')
            if j > 0 and not ((j - 1) >= 6 and (j - 1) - i >= 2) and not (i >= 6 and i - (j - 1) >= 2) and not (j - 1 == 6 and i < 5):
                t.append(f'{cell(i, j-1)}*(1-{a_game(i + j - 1 + 1)})')
            ws.cell(r0 + i, c0 + j, '=' + ('+'.join(t) if t else '0'))
    terms = [cell(6, j) for j in range(5)] + [f'{cell(6, 5)}*{a_game(12)}', f'{cell(6, 6)}*{tbc}']
    ws.cell(r0 + 7, c0 - 1, 'P(A wins set)')
    ws.cell(r0 + 7, c0, '=' + '+'.join(terms))
    return ws.cell(r0 + 7, c0).coordinate

def model_sheet(wb, key):
    r = res[key]; A, B = r['A'], r['B']; ws = wb.create_sheet(f'Model {key}')
    ws['A1'] = f'{key}: {A} vs {B} - editable model (blue = input; everything else is a live formula)'; ws['A1'].font = TITLE
    phys = key == 'QF-B'
    inputs = [('pa_ratings', f'{A} SPW vs {B} from ridge (as of 2026-10-07)', r['pa_ratings']),
              ('pb_ratings', f'{B} SPW vs {A} from ridge', r['pb_ratings']),
              ('pace', 'Court pace (SPW, both players)', 0.010),
              ('fsA', f'{A} this-week serve residual (raw)', r['formA_serve']), ('frA', f'{A} this-week return residual (raw)', r['formA_return']),
              ('fsB', f'{B} this-week serve residual (raw)', r['formB_serve']), ('frB', f'{B} this-week return residual (raw)', r['formB_return']),
              ('w', 'Form weight (backtest best 0.10; spec 0.50)', 0.10),
              ('hwA', 'H2H weight, A serving = n/(n+1000)', r.get('h2h_wA', 0.0)), ('hA', 'H2H SPW, A serving (time-weighted)', r.get('h2h_spwA', 0.0)),
              ('hwB', 'H2H weight, B serving', r.get('h2h_wB', 0.0)), ('hB', 'H2H SPW, B serving', r.get('h2h_spwB', 0.0)),
              ('plim', 'P(limited) physical state (Zverev ankle)' if phys else 'P(limited) (n/a: 0)', 0.50 if phys else 0.0),
              ('sh1', 'Limited: serve hit set 1', 0.005 if phys else 0.0), ('sh2', 'Limited: serve hit set 2', 0.020 if phys else 0.0), ('sh3', 'Limited: serve hit set 3', 0.030 if phys else 0.0),
              ('rh1', 'Limited: opponent SPW up set 1', 0.005 if phys else 0.0), ('rh2', 'Limited: opponent SPW up set 2', 0.015 if phys else 0.0), ('rh3', 'Limited: opponent SPW up set 3', 0.025 if phys else 0.0),
              ('limA', 'Limited player is A? (1 = A, 0 = B)', 1),
              ('f3', 'Set-3 fatigue: opponent SPW up in set 3 (ADF return, QF-A)', 0.010 if key == 'QF-A' else 0.0),
              ('f3A', 'Fatigued player is B? (1 = B tired -> A SPW up)', 1),
              ('k', 'Winner shrink k (walk-forward backtest)', 0.58),
              ('mkt', 'Market consensus P(A) no-vig (edit)', mk[key]['consensus']),
              ('gem', 'Gemini blind P(A)', REC[key]['blind']['Gemini']), ('grok', 'Grok blind P(A)', REC[key]['blind']['Grok']),
              ('wC', 'Pool weight Claude', REC[key]['weights']['Claude']), ('wG', 'Pool weight Gemini', REC[key]['weights']['Gemini']), ('wK', 'Pool weight Grok', REC[key]['weights']['Grok'])]
    R = {}
    ws['A3'] = 'INPUTS'; ws['A3'].font = HDR
    for i, (k, lab, v) in enumerate(inputs):
        ws.cell(4 + i, 1, lab); c = ws.cell(4 + i, 2, float(v)); c.font = BLUE; c.fill = FILL; R[k] = f'$B${4 + i}'
    r0 = 4 + len(inputs) + 1
    ws.cell(r0, 1, 'CALCULATIONS').font = HDR
    calc = [('pa_fit', 'A SPW fit state', f"={R['pa_ratings']}+{R['pace']}+{R['w']}*({R['fsA']}-{R['frB']})"),
            ('pb_fit', 'B SPW fit state', f"={R['pb_ratings']}+{R['pace']}+{R['w']}*({R['fsB']}-{R['frA']})")]
    for i, (k, lab, f) in enumerate(calc):
        ws.cell(r0 + 1 + i, 1, lab); ws.cell(r0 + 1 + i, 2, f); R[k] = f'$B${r0 + 1 + i}'
    rr = r0 + 3
    ws.cell(rr, 1, 'A SPW after H2H blend'); ws.cell(rr, 2, f"=(1-{R['hwA']})*{R['pa_fit']}+{R['hwA']}*({R['hA']}+{R['pace']})*({R['hwA']}>0)+{R['hwA']}*{R['pa_fit']}*({R['hwA']}=0)"); R['pa'] = f'$B${rr}'
    ws.cell(rr + 1, 1, 'B SPW after H2H blend'); ws.cell(rr + 1, 2, f"=(1-{R['hwB']})*{R['pb_fit']}+{R['hwB']}*({R['hB']}+{R['pace']})*({R['hwB']}>0)+{R['hwB']}*{R['pb_fit']}*({R['hwB']}=0)"); R['pb'] = f'$B${rr + 1}'
    # per-set SPW in fit and limited states
    rr += 3; ws.cell(rr, 1, 'Per-set SPW').font = HDR; ws.cell(rr, 2, 'Set 1'); ws.cell(rr, 3, 'Set 2'); ws.cell(rr, 4, 'Set 3')
    states = {}
    for si, (st, lim) in enumerate((('fit', 0), ('limited', 1))):
        for pi, who in enumerate(('A', 'B')):
            row = rr + 1 + si * 2 + pi; ws.cell(row, 1, f'{who} SPW ({st})')
            for s in range(3):
                sh = R[f'sh{s+1}']; rh = R[f'rh{s+1}']
                base = R['pa'] if who == 'A' else R['pb']
                if who == 'A':
                    f = f"={base}-{lim}*({R['limA']}*{sh}-(1-{R['limA']})*{rh})"
                else:
                    f = f"={base}-{lim}*((1-{R['limA']})*{sh}-{R['limA']}*{rh})"
                if s == 2:
                    f += f"+{R['f3']}*({R['f3A']}*{int(who == 'A')}+(1-{R['f3A']})*{int(who == 'B')})"
                ws.cell(row, 2 + s, f)
            states[(st, who)] = [ws.cell(row, 2 + s).coordinate for s in range(3)]
    # DP blocks: for each state and set, TB + set grids
    col = 7; ws.cell(2, col, 'DP grids (rows = A points/games, cols = B) - live formulas').font = HDR
    setwin = {}
    top = 4
    for st in ('fit', 'limited'):
        for s in range(3):
            pa_c = '$' + states[(st, 'A')][s][0] + '$' + states[(st, 'A')][s][1:]; pb_c = '$' + states[(st, 'B')][s][0] + '$' + states[(st, 'B')][s][1:]
            ws.cell(top, col, f'{st} set {s+1}: TB (A serves first)').font = HDR
            tbc = tb_block(ws, top + 1, col + 1, pa_c, pb_c)
            ws.cell(top, col + 9, f'{st} set {s+1}: set (A serves first)').font = HDR
            hA = hold_f(pa_c); hB = hold_f(pb_c)
            sc = set_block(ws, top + 1, col + 10, hA, hB, tbc)
            setwin[(st, s)] = sc; top += 10
    rr = rr + 6; ws.cell(rr, 1, 'RESULTS').font = HDR
    out = []
    for st in ('fit', 'limited'):
        s1, s2, s3 = (setwin[(st, s)] for s in range(3))
        out.append((f'P(A wins match), {st} state (iid)', f'={s1}*{s2}+{s1}*(1-{s2})*{s3}+(1-{s1})*{s2}*{s3}'))
    for i, (lab, f) in enumerate(out):
        ws.cell(rr + 1 + i, 1, lab); ws.cell(rr + 1 + i, 2, f)
    pf, pl = f'$B${rr + 1}', f'$B${rr + 2}'
    rows2 = [('P(A) iid mixture', f"=(1-{R['plim']})*{pf}+{R['plim']}*{pl}"),
             ('P(A) BLIND = shrink(mixture)', f"=1/(1+EXP(-{R['k']}*LN(B{rr+3}/(1-B{rr+3}))))"),
             ('P(A) POOLED blind (Claude/Gemini/Grok, log-odds)', f"=1/(1+EXP(-({R['wC']}*LN(B{rr+4}/(1-B{rr+4}))+{R['wG']}*LN({R['gem']}/(1-{R['gem']}))+{R['wK']}*LN({R['grok']}/(1-{R['grok']})))))"),
             ('P(A) FINAL = 50/50 log-odds blend of pooled blind with market', f"=1/(1+EXP(-(0.5*LN(B{rr+5}/(1-B{rr+5}))+0.5*LN({R['mkt']}/(1-{R['mkt']})))))"),
             ('Fair American odds A (final)', f'=IF(B{rr+6}>=0.5,-100*B{rr+6}/(1-B{rr+6}),100*(1-B{rr+6})/B{rr+6})'),
             ('Python check: blind P(A) (exact engine)', r['p_win']), ('Python check: final P(A)', mk[key]['final'])]
    for i, (lab, f) in enumerate(rows2):
        ws.cell(rr + 3 + i, 1, lab); ws.cell(rr + 3 + i, 2, f)
    ws.column_dimensions['A'].width = 58; ws.column_dimensions['B'].width = 14
    return ws, rr

def pricing_sheet(wb):
    ws = wb.create_sheet('Pricing')
    ws['A1'] = 'Price ladder - enter the CURRENT DraftKings price in column F (blue). EV, break-even and 1/4 Kelly (cap 2%) recalculate.'; ws['A1'].font = TITLE
    hdr = ['Match', 'Market', 'Selection', 'Model P', 'Fair odds', 'DK price (enter)', 'Price source', 'Break-even', 'EV per $1', '1/4 Kelly (cap 2%)', 'Take at (EV>=2.5%)', 'Bet?']
    for j, h in enumerate(hdr): ws.cell(3, 1 + j, h).font = HDR
    r = 4
    for k in ('QF-A', 'QF-B'):
        for m in card[k]['markets']:
            ws.cell(r, 1, k); ws.cell(r, 2, m['market']); ws.cell(r, 3, m['selection']); ws.cell(r, 4, round(m['p'], 4))
            ws.cell(r, 5, f'=IF(D{r}>=0.5,-100*D{r}/(1-D{r}),100*(1-D{r})/D{r})')
            c = ws.cell(r, 6, m.get('price')); c.font = BLUE; c.fill = FILL
            ws.cell(r, 7, m.get('price_src', 'enter DK price'))
            ws.cell(r, 8, f'=IF(F{r}="","",IF(F{r}>0,100/(F{r}+100),-F{r}/(-F{r}+100)))')
            ws.cell(r, 9, f'=IF(F{r}="","",D{r}*IF(F{r}>0,1+F{r}/100,1+100/-F{r})-1)')
            ws.cell(r, 10, f'=IF(F{r}="","",MAX(0,MIN(0.02,0.25*((IF(F{r}>0,F{r}/100,100/-F{r}))*D{r}-(1-D{r}))/IF(F{r}>0,F{r}/100,100/-F{r}))))')
            ws.cell(r, 11, m['take_at'])
            ws.cell(r, 12, f'=IF(F{r}="","",IF(I{r}>=0.025,"YES","no"))')
            r += 1
    r += 1; ws.cell(r, 1, 'Same-game parlays (joint probability from the 60k simulation; never the naive product)').font = HDR; r += 1
    for j, h in enumerate(['Match', 'SGP', 'Joint P', 'Naive product', 'Fair odds', 'DK SGP price (enter)', 'EV per $1', 'Take at (EV>=5%)']): ws.cell(r, 1 + j, h).font = HDR
    r += 1
    for k in ('QF-A', 'QF-B'):
        for s in card[k]['sgps']:
            ws.cell(r, 1, k); ws.cell(r, 2, s['sgp']); ws.cell(r, 3, round(s['p'], 4)); ws.cell(r, 4, round(s['naive_product'], 4))
            ws.cell(r, 5, f'=IF(C{r}>=0.5,-100*C{r}/(1-C{r}),100*(1-C{r})/C{r})'); c = ws.cell(r, 6, None); c.font = BLUE; c.fill = FILL
            ws.cell(r, 7, f'=IF(F{r}="","",C{r}*IF(F{r}>0,1+F{r}/100,1+100/-F{r})-1)'); ws.cell(r, 8, s['take_at']); r += 1
    for j, w in enumerate([7, 20, 42, 9, 10, 14, 34, 10, 10, 14, 14, 6]): ws.column_dimensions[get_column_letter(1 + j)].width = w

def table_sheet(wb, name, df, note=None):
    ws = wb.create_sheet(name[:31]); r0 = 1
    if note: ws.cell(1, 1, note).font = HDR; r0 = 3
    for j, c in enumerate(df.columns): ws.cell(r0, 1 + j, str(c)).font = HDR
    for i, row in enumerate(df.itertuples(index=False)):
        for j, v in enumerate(row):
            if isinstance(v, (np.floating, float)): v = None if np.isnan(v) else float(v)
            elif isinstance(v, (np.integer,)): v = int(v)
            elif not isinstance(v, (int, str, type(None))): v = str(v)
            ws.cell(r0 + 1 + i, 1 + j, v)
    for j in range(len(df.columns)): ws.column_dimensions[get_column_letter(1 + j)].width = max(10, min(45, len(str(df.columns[j])) + 4))

def build(path):
    wb = Workbook(); ws = wb.active; ws.title = 'README'
    lines = ['Deep Sports Engine - Shanghai 2026 QF model workbook (built 2026-10-15)',
             'Model tabs: blue cells are inputs; every other number is a live formula (hold, 7-pt tiebreak DP, set DP, Bo3, mixture, shrink, blend).',
             'Pricing tab: type the current DraftKings price into column F; EV, break-even and quarter-Kelly recalculate.',
             'Raw tabs: walk-forward backtest (1,667 hard Bo3 tour-level matches, 2025-01-06..2026-09-30), calibration, form, H2H, charting analytics, markets.',
             'Data: TML database (2023-01..2026-01, exact counts), Tennis Abstract 2026 tournament tables via Exa (17 hard events; serve points estimated from score),',
             '      Shanghai 2026 match stats (tennis.com / ATP Infosys, DeepCourt), Match Charting Project points (last charted 2026-05-21).',
             'Model probabilities in Pricing come from the re-centred 60,000-sim Monte Carlo (day-form sd 0.06) and are values, not formulas.',
             '21+, 1-800-GAMBLER; not betting advice.']
    for i, l in enumerate(lines): ws.cell(1 + i, 1, l).font = TITLE if i == 0 else Font()
    ws.column_dimensions['A'].width = 140
    for k in ('QF-A', 'QF-B'): model_sheet(wb, k)
    rr = []
    for k in ('QF-A', 'QF-B'):
        for a in ('Claude', 'Gemini', 'Grok'):
            g = REC['grades'][a]
            rr.append(dict(match=k, ai=a, blind_pA=REC[k]['blind'][a], weight=REC[k]['weights'][a], facts=g['facts'], opponent_adjustment=g['opponent_adjustment'],
                           calibration=g['calibration'], real_prices=g['real_prices'], ev=g['ev'], audit_notes=g['notes']))
        rr.append(dict(match=k, ai='POOLED blind', blind_pA=REC[k]['pooled_blind'], weight=1.0, audit_notes=f"consensus {REC[k]['consensus']:.4f} -> FINAL {REC[k]['final_pooled']:.4f} (pre-reconciliation {REC[k]['final_pre_recon']:.4f})"))
    recon_df = pd.DataFrame(rr)
    pricing_sheet(wb)
    table_sheet(wb, 'Reconciliation', recon_df, 'Step 9: grades (0-10), pool weights (start 0.6/0.2/0.2), pooled blind and final; P(A) = Gea (QF-A), Zverev (QF-B)')
    table_sheet(wb, 'Other AI picks priced', pd.DataFrame(XC), "Every AI's picks priced on the same re-centred 60k simulation; prices for Gemini derivatives are its own [INFERENCE] numbers, not DK quotes")
    table_sheet(wb, 'Backtest variants', pd.read_csv(f'{OUT}/bt_variants.csv'), 'Walk-forward backtest, n=1667 hard Bo3 tour-level, 2025-01-06..2026-09-30')
    table_sheet(wb, 'Backtest form-H2H', pd.read_csv(f'{OUT}/bt_form_h2h.csv'), 'Form weight x H2H blend (day-form sd 0.085)')
    table_sheet(wb, 'Calib derivatives sd06', pd.read_csv(f'{OUT}/bt_deriv_calib_sd06.csv'), 'Sets / tiebreak / games calibration at day-form sd 0.06')
    table_sheet(wb, 'Calib derivatives sd085', pd.read_csv(f'{OUT}/bt_deriv_calib.csv'), 'Same at sd 0.085 (your Lesson-12 setting)')
    table_sheet(wb, 'Calib winner deciles', pd.read_csv(f'{OUT}/bt_calib_deciles.csv'), 'Winner calibration by decile (favourite side)')
    table_sheet(wb, 'Charting analytics', pd.read_csv(f'{OUT}/mcp_analytics.csv').rename(columns={'Unnamed: 0': 'stat'}), 'Match Charting Project, hard, 2024-05-21..2026-05-21 (Gea: 0 charted points)')
    ft = pd.DataFrame(res['QF-A']['form_table'] + res['QF-B']['form_table'])
    table_sheet(wb, 'Shanghai week form', ft, 'Opponent-adjusted serve/return residuals this week (expectation includes +1.0 pt pace)')
    table_sheet(wb, 'Zverev vs elite returners', pd.read_csv(f'{OUT}/zverev_vs_elite_returners.csv'), 'Zverev hard-court SPW vs expectation, 2024-10..2026-10 (elite = top-12 return ratings)')
    from qf_model import H2H_ZA
    table_sheet(wb, 'H2H Zverev-de Minaur', pd.DataFrame(H2H_ZA, columns=['date', 'event', 'Z_svpt', 'Z_SPW', 'ADM_svpt', 'ADM_SPW', 'include_frac (None = match-TB points removed)']), 'Hard-court H2H serve data (Tennis Abstract); RG 2024 clay excluded')
    mrows = []
    for k in ('QF-A', 'QF-B'):
        for m in mk[k]['market_table']: mrows.append(dict(match=k, **m))
    table_sheet(wb, 'Markets', pd.DataFrame(mrows), 'Market numbers (no-vig P for player A); status column says what was opened vs search-reported')
    wrows = []
    for k in ('QF-A', 'QF-B'):
        for s, v in mk[k]['waterfall']: wrows.append(dict(match=k, step=s, p_A=v))
        inv = mk[k]['inversion']; wrows.append(dict(match=k, step=f"market-implied serve probs pa={inv['pa']:.4f} pb={inv['pb']:.4f} (sum {inv['sum']:.3f} vs model {inv['model_sum']:.3f})", p_A=inv['p_win']))
    table_sheet(wb, 'Waterfall + inversion', pd.DataFrame(wrows), 'Blind-number decomposition and market inversion')
    sims = []
    for k in ('QF-A', 'QF-B'):
        for g, c in enumerate(card[k]['games_hist']): sims.append(dict(match=k, total_games=g, sims=c))
    table_sheet(wb, 'Total games dist', pd.DataFrame(sims), 'Re-centred simulation: total-games histogram (60,000 sims per match)')
    wb.save(path)

if __name__ == '__main__':
    p = os.path.join(OUT, 'Shanghai_2026_QF_model.xlsx'); build(p); print('saved', p)
